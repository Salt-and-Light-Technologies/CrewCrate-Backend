"""Bounded spreadsheet decoding; no formulas are evaluated or original files stored."""
import csv
import io
from pathlib import PurePath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .imports import MAX_BYTES, MAX_ROWS, parse_csv


def headers(row):
    values = [str(x or "").strip() for x in row]
    if (not values or len(values) > 128 or any(not x or len(x) > 200 for x in values)
            or len({x.casefold() for x in values}) != len(values)):
        raise HTTPException(422, "Use a first row with unique, nonempty column names (up to 128 columns).")
    return values


def decode_sheet(raw, filename, sheet=None):
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "File exceeds 5 MB")
    extension = PurePath(filename or "").suffix.lower()
    if extension not in (".csv", ".tsv", ".xlsx"):
        raise HTTPException(422, "Upload CSV, TSV or Excel .xlsx. Export Numbers, Google Sheets or older Excel files first.")
    workbook = None
    try:
        if extension != ".xlsx":
            text = raw.decode("utf-8-sig")
            delimiter = "\t" if extension == ".tsv" else ","
            if extension == ".csv":
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
                except csv.Error:
                    pass
            rows = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            first = headers(next(rows))
            sheet_names = []
            chosen = None
            workbook = None
        else:
            with ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 500 or sum(x.file_size for x in archive.infolist()) > 20_000_000:
                    raise HTTPException(413, "Expanded Excel workbook is too large")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheet_names = [ws.title for ws in workbook.worksheets]
            if not sheet_names:
                raise HTTPException(422, "Workbook has no contact worksheets")
            chosen = sheet or sheet_names[0]
            if chosen not in sheet_names:
                workbook.close()
                raise HTTPException(422, "Choose a worksheet from this file")
            ws = workbook[chosen]
            if (ws.max_column or 0) > 128 or (ws.max_row or 0) > MAX_ROWS + 1:
                workbook.close()
                raise HTTPException(413, "Worksheet exceeds 128 columns or 20,000 rows")
            rows = ws.iter_rows(values_only=True, max_row=min(ws.max_row or MAX_ROWS + 2, MAX_ROWS + 2), max_col=min(ws.max_column or 128, 128))
            header = list(next(rows))
            while header and header[-1] is None:
                header.pop()
            first = headers(header)
        try:
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(first)
            count = 0
            for index, row in enumerate(rows):
                if index >= MAX_ROWS:
                    raise HTTPException(413, "Worksheet exceeds 20,000 rows")
                if extension == ".xlsx":
                    if any(x is not None for x in row[len(first):]):
                        raise HTTPException(422, "Every populated column needs a header")
                    row = row[:len(first)]
                values = ["" if x is None else str(int(x)) if isinstance(x, float) and x.is_integer() else str(x) for x in row]
                if not any(x.strip() for x in values):
                    continue
                count += 1
                if count > MAX_ROWS or output.tell() > MAX_BYTES:
                    raise HTTPException(413, "Worksheet exceeds 20,000 rows or 5 MB of contact data")
                if extension == ".xlsx" and any(x.startswith("=") for x in values):
                    raise HTTPException(422, "Use plain values instead of formulas in your contact worksheet")
                writer.writerow(values)
            normalized = output.getvalue().encode("utf-8")
            if len(normalized) > MAX_BYTES:
                raise HTTPException(413, "Contact data exceeds 5 MB")
            return normalized, {"columns": first, "sheets": sheet_names, "selected_sheet": chosen}
        finally:
            if workbook:
                workbook.close()
    except HTTPException:
        raise
    except (UnicodeError, csv.Error, StopIteration, BadZipFile, ValueError, KeyError, TypeError, OSError, IndexError, AttributeError, RuntimeError, OverflowError, ParseError, DefusedXmlException, InvalidFileException):
        raise HTTPException(422, "Couldn’t read this file. Use UTF-8 CSV/TSV or an unencrypted Excel .xlsx with headers in row 1.") from None
    finally:
        if workbook:
            workbook.close()


def parse_file(raw, filename, sheet, phone_column, name_column, email_column, permission_columns="{}", permission_attested=False):
    normalized, _ = decode_sheet(raw, filename, sheet)
    return parse_csv(normalized, phone_column, name_column, email_column, permission_columns, permission_attested)
